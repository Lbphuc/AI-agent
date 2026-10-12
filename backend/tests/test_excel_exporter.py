"""Tests cho module excel_exporter và endpoint /export/xlsx."""
from __future__ import annotations

import io
from datetime import datetime, timezone

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.modules.excel_exporter import generate_excel_report, safe_sheet_name
from backend.app.schemas import (
    AnalysisGoal,
    AnalysisTask,
    AnalyzeResponse,
    ChartSpec,
    ExecutionResult,
    ExportSections,
    ExportXlsxRequest,
    InsightResult,
    ReportMeta,
    TaskResult,
)


@pytest.fixture
def sample_analyze_response() -> AnalyzeResponse:
    goal = AnalysisGoal(
        question="So sánh doanh thu giữa các chi nhánh",
        metric="Doanh_thu_trieu_vnd",
        group_by="Chi_nhanh",
        tasks=[AnalysisTask(type="comparison", metric="Doanh_thu_trieu_vnd", group_by="Chi_nhanh")],
    )
    chart1 = ChartSpec(
        chart_type="bar",
        title="So sánh doanh thu theo chi nhánh",
        x=["Hà Nội", "TP.HCM", "Đà Nẵng", "=1+1", "@calc"],  # Bao gồm nhãn thử formula injection
        series={"Doanh thu": [2500.0, 3200.0, 1800.0, 1500.0, 900.0]},
        role="primary",
        caption="TP.HCM dẫn đầu doanh thu với 3.200 triệu VNĐ.",
        unit="triệu VNĐ",
    )
    chart2 = ChartSpec(
        chart_type="pie",
        title="Tỷ trọng doanh thu theo chi nhánh",
        x=["Hà Nội", "TP.HCM", "Đà Nẵng"],
        series={"Tỷ trọng (%)": [33.3, 42.7, 24.0]},
        role="supplementary",
        perspective="share_of_total",
        caption="TP.HCM chiếm tỷ trọng lớn nhất (42,7%).",
        unit="%",
    )
    meta = ReportMeta(
        generated_at=datetime.now(timezone.utc).isoformat(),
        model_name="Gemini-2.0-Flash",
        rows_total=100,
        rows_used=80,
        goal_summary="So sánh doanh thu giữa các chi nhánh.",
        notes=["Dữ liệu mẫu tuần 7"],
    )
    insight = InsightResult(
        insight_text="TP.HCM là chi nhánh có doanh thu cao nhất (3.200 triệu VNĐ), tiếp theo là Hà Nội.",
        accuracy_score=1.0,
        known_values=[3200.0, 2500.0, 1800.0],
    )
    return AnalyzeResponse(
        goal=goal,
        task_results=[TaskResult(task=goal.tasks[0], execution=ExecutionResult(success=True, results={}))],
        charts=[chart1, chart2],
        insight=insight,
        meta=meta,
    )


def test_safe_sheet_name_sanitization():
    existing: set[str] = set()
    s1 = safe_sheet_name("Tỷ trọng [2026]: Chi nhánh / TP * HN", existing)
    assert not any(c in s1 for c in "[]:*?/\\")
    assert len(s1) <= 31
    assert s1 in existing

    # Duplicate names get numbered
    s2 = safe_sheet_name("Tỷ trọng [2026]: Chi nhánh / TP * HN", existing)
    assert s2 != s1
    assert s2.endswith("_1")
    assert s2 in existing


def test_generate_excel_report_content_and_openpyxl_validation(sample_analyze_response: AnalyzeResponse):
    req = ExportXlsxRequest(
        response=sample_analyze_response,
        sections=ExportSections(),
        org_name="Công ty Cổ phần ABC",
    )
    raw_bytes = generate_excel_report(req)
    assert raw_bytes.startswith(b"PK")  # ZIP / OpenXML header

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    sheet_names = wb.sheetnames

    # Check sheets exist
    assert "Tóm tắt" in sheet_names
    assert "Cách tính & Lưu ý" in sheet_names
    assert len(sheet_names) >= 4  # Tóm tắt + 2 charts + Cách tính

    # Check Summary sheet
    ws_summary = wb["Tóm tắt"]
    summary_text = [cell.value for row in ws_summary.iter_rows() for cell in row if cell.value]
    assert any("Công ty Cổ phần ABC" in str(v) for v in summary_text)
    assert any("So sánh doanh thu giữa các chi nhánh" in str(v) for v in summary_text)

    # Check Chart 1 sheet: formula injection prevention
    ws_chart1 = wb[sheet_names[1]]
    col_x_values = [cell.value for cell in ws_chart1["A"] if cell.value]

    # "=1+1" must be saved as text, cell data_type must be string 's'
    formula_cell = None
    for row in ws_chart1.iter_rows():
        for cell in row:
            if cell.value == "=1+1":
                formula_cell = cell
                break
    assert formula_cell is not None
    assert formula_cell.data_type == "s"  # string, NOT formula 'f'

    # Check numeric cell type
    num_cell = None
    for row in ws_chart1.iter_rows():
        for cell in row:
            if cell.value == 3200.0 or cell.value == 2500.0:
                num_cell = cell
                break
    assert num_cell is not None
    assert num_cell.data_type in ("n", "f")  # numeric


def test_export_xlsx_endpoint_success(sample_analyze_response: AnalyzeResponse):
    client = TestClient(app)
    req_data = {
        "response": sample_analyze_response.model_dump(),
        "sections": {
            "insight": True,
            "primary_chart": True,
            "supplementary_charts": True,
            "data_tables": True,
            "calculation_notes": True,
        },
        "org_name": "Tập đoàn XYZ",
    }
    resp = client.post("/export/xlsx", json=req_data)
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "Content-Disposition" in resp.headers
    assert "BaoCao_PhanTich" in resp.headers["Content-Disposition"]
    assert resp.content.startswith(b"PK")


def test_export_xlsx_endpoint_validation_error():
    client = TestClient(app)
    resp = client.post("/export/xlsx", json={"invalid": "payload"})
    assert resp.status_code == 422
